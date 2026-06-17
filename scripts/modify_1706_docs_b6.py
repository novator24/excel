from pathlib import Path
import re

from docx import Document


BASE = Path(r"C:\Users\leto\Documents\GITLAB\excel")
OUT_DIR = BASE / "test" / "1706"
TMP_DIR = OUT_DIR / "_tmp"


def split_sentences(text: str):
    return [m.group(0) for m in re.finditer(r"[^.!?]+[.!?]|[^.!?]+$", text) if m.group(0).strip()]


def lower_first_alpha(text: str) -> str:
    chars = list(text)
    for i, ch in enumerate(chars):
        if ch.isalpha():
            chars[i] = ch.lower()
            break
    return "".join(chars)


def bump_plus_pattern(text: str) -> str:
    # Example: "10+ years" -> "13 years", "2+ months" -> "5 months"
    return re.sub(r"(\d+)\+\s", lambda m: f"{int(m.group(1)) + 3} ", text)


def modify_doc(src: Path, dst: Path, replacement: str) -> None:
    doc = Document(str(src))

    # Change all N+ <space> patterns across all paragraphs
    for p in doc.paragraphs:
        p.text = bump_plus_pattern(p.text)

    sentences = []
    for p_idx, paragraph in enumerate(doc.paragraphs):
        for sentence in split_sentences(paragraph.text):
            sentences.append((p_idx, sentence, len(sentence.strip())))

    if not sentences:
        doc.save(str(dst))
        return

    longest = max(sentences, key=lambda x: x[2])
    shortest = min(sentences, key=lambda x: x[2])

    # 1) Longest sentence: replace first space
    p_long = doc.paragraphs[longest[0]]
    long_text = longest[1]
    idx = long_text.find(" ")
    if idx != -1:
        long_mod = long_text[:idx] + replacement + long_text[idx + 1 :]
        p_long.text = p_long.text.replace(long_text, long_mod, 1)

    # 2) Shortest sentence: lowercase first alphabetic char
    p_short = doc.paragraphs[shortest[0]]
    if shortest[0] == longest[0]:
        candidates = split_sentences(p_short.text)
        short_text = min(candidates, key=lambda s: len(s.strip())) if candidates else shortest[1]
    else:
        short_text = shortest[1]
    p_short.text = p_short.text.replace(short_text, lower_first_alpha(short_text), 1)

    doc.save(str(dst))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)

    modify_doc(BASE / "CV_BASOVI.docx", TMP_DIR / "CV_BASOVI_mod_B6.docx", "_")
    modify_doc(BASE / "RESUME_BASOVI.docx", TMP_DIR / "RESUME_BASOVI_mod_B6.docx", ")")
    modify_doc(BASE / "LETTER_BASOVI.docx", TMP_DIR / "LETTER_BASOVI_mod_B6.docx", "'")
    print("MOD_B6_OK")


if __name__ == "__main__":
    main()
