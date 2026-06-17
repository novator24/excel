from pathlib import Path
import re

from docx import Document


BASE = Path(r"C:\Users\leto\Documents\GITLAB\excel")
OUT_DIR = BASE / "test" / "1706"
TMP_DIR = OUT_DIR / "_tmp"


def split_sentences(text: str):
    return [m.group(0) for m in re.finditer(r"[^.!?]+[.!?]|[^.!?]+$", text) if m.group(0).strip()]


def lower_first_alpha(text: str) -> str:
    chars = list(text)
    for i, ch in enumerate(chars):
        if ch.isalpha():
            chars[i] = ch.lower()
            break
    return "".join(chars)


def modify_docx(src: Path, dst: Path, replace_char: str) -> None:
    doc = Document(str(src))

    sentence_refs = []
    for p_idx, p in enumerate(doc.paragraphs):
        for sent in split_sentences(p.text):
            sentence_refs.append((p_idx, sent, len(sent.strip())))

    if not sentence_refs:
        doc.save(str(dst))
        return

    longest = max(sentence_refs, key=lambda x: x[2])
    shortest = min(sentence_refs, key=lambda x: x[2])

    # Edit longest sentence: replace first space with required character
    p_long = doc.paragraphs[longest[0]]
    long_text = longest[1]
    first_space = long_text.find(" ")
    if first_space != -1:
        long_text_mod = long_text[:first_space] + replace_char + long_text[first_space + 1 :]
        p_long.text = p_long.text.replace(longest[1], long_text_mod, 1)

    # Edit shortest sentence: first alphabetic letter to lowercase
    p_short = doc.paragraphs[shortest[0]]
    if shortest[0] == longest[0]:
        # Re-evaluate in updated paragraph
        updated_sents = split_sentences(p_short.text)
        if updated_sents:
            shortest_text = min(updated_sents, key=lambda s: len(s.strip()))
        else:
            shortest_text = shortest[1]
    else:
        shortest_text = shortest[1]
    p_short.text = p_short.text.replace(shortest_text, lower_first_alpha(shortest_text), 1)

    doc.save(str(dst))


def export_pdf(src_docx: Path, dst_pdf: Path) -> None:
    import win32com.client  # type: ignore

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    wd_format_pdf = 17
    try:
        doc = word.Documents.Open(str(src_docx))
        try:
            doc.SaveAs(str(dst_pdf), FileFormat=wd_format_pdf)
        finally:
            doc.Close(False)
    finally:
        word.Quit()


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)

    cv_src = BASE / "CV_BASOVI.docx"
    resume_src = BASE / "RESUME_BASOVI.docx"
    letter_src = BASE / "LETTER_BASOVI.docx"
    welcome_src = BASE / "WELCOME_BASOVI.docx"

    cv_mod = TMP_DIR / "CV_BASOVI_mod.docx"
    resume_mod = TMP_DIR / "RESUME_BASOVI_mod.docx"
    letter_mod = TMP_DIR / "LETTER_BASOVI_mod.docx"

    modify_docx(cv_src, cv_mod, "_")
    modify_docx(resume_src, resume_mod, ")")
    modify_docx(letter_src, letter_mod, "'")

    export_pdf(cv_mod, OUT_DIR / "CV_BASOVI_JD.pdf")
    export_pdf(welcome_src, OUT_DIR / "WELCOME_BASOVI_JD.pdf")
    export_pdf(resume_mod, OUT_DIR / "RESUME_BASOVI_JD.pdf")
    export_pdf(letter_mod, OUT_DIR / "LETTER_BASOVI_JD.pdf")

    print("EXPORT_OK")
    print("CV_BASOVI_JD.pdf")
    print("WELCOME_BASOVI_JD.pdf")
    print("RESUME_BASOVI_JD.pdf")
    print("LETTER_BASOVI_JD.pdf")


if __name__ == "__main__":
    main()
